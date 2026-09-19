"""Sheet Tồn kho của số dư ban đầu (nhóm 5, lát 8C-1 — RT-24, FR-OPB-004/007).

Lời hứa của lát:

* mỗi dòng sheet là một LẦN NHẬP → một lớp `opening_balance_stock_layers` → một
  movement nhập đã có giá, ngày = ngày đầu năm − 1, kỳ = kỳ đầu năm, thứ tự
  trong ngày theo ngày nhập (FIFO ăn đúng lớp cũ trước);
* số lượng gõ theo đơn vị phụ quy về đơn vị chính (FR-STK-006); dòng cha nhóm 5
  đổ vào cân đối tài khoản kỳ 1 như mọi nhóm khác;
* TK kho theo dõi chiều kho/mã hàng phải nhập ở sheet này; TK không phải TK kho,
  mã hàng dịch vụ, đơn vị chưa khai, sổ quản trị, năm FIFO thiếu ngày nhập →
  lỗi đúng mã;
* thay-trọn nhóm 5 gỡ movement cũ và đánh dấu tính lại khi năm đã có xuất; xóa
  nhóm cũng gỡ movement; chi nhánh đã có sổ kho trước đầu năm → từ chối nhập tay.
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from import_support import FakeProgress
from inventory_support import (
    BOX_FACTOR,
    MAIN_WAREHOUSE_ID,
    UNIT_BOX_ID,
    create_goods_item,
    issue_payload,
    marks_of_branch,
    out_movement,
    post_issue,
    post_receipt,
    receipt_payload,
    run_engine,
    seed_inventory_package_data,
    set_valuation_method,
)
from ket.kernel.datasets.provisioning import DatasetRef
from ket.kernel.errors import (
    InventoryLayerReferencedError,
    InventoryMovementBeforeOpeningStockError,
    OpeningStockHistoryExistsError,
)
from ket.kernel.jobs.registry import JobContext
from ket.kernel.master_data.models.item_unit import ItemUnit
from ket.kernel.periods.models import FiscalYear, InventoryValuationMethod
from ket.kernel.persistence.unit_of_work import unit_of_work
from ket.modules.inventory.models import CostState, InventoryMovement
from ket.modules.inventory.service import InventoryVoucherService
from ket.modules.inventory.stock import stock_rows
from ket.posting.balances.query_service import trial_balance
from ket.posting.engine.models import Ledger
from ket.posting.opening_balances.clear_job import ClearGroupParams, run_clear_group
from ket.posting.opening_balances.models import (
    OpeningBalance,
    OpeningBalanceStockLayer,
    OpeningDetailKind,
)
from ket.posting.opening_balances.template import (
    ACCOUNT_SHEET,
    STOCK_SHEET,
    build_opening_template,
)
from posting_support import PostingContext, posting_scope, seed_posting_context
from test_opening_balance_parsing import workbook_of
from test_opening_balances_carry_forward import carry_forward, ensure_fiscal_year
from test_opening_balances_import import first_period_id, run_opening_import

pytestmark = pytest.mark.db

ACTOR_ID = 1
DEC_31 = date(2025, 12, 31)
NOV_3, DEC_1 = date(2025, 11, 3), date(2025, 12, 1)
JAN_10, JAN_20 = date(2026, 1, 10), date(2026, 1, 20)
MAIN_CODE, SECOND_CODE = "KHO-8A-1", "KHO-8A-2"


@pytest.fixture
def context(session_factory: sessionmaker[Session], dataset_alpha: DatasetRef) -> PostingContext:
    return seed_posting_context(session_factory, dataset_alpha)


@pytest.fixture
def accounts(
    session_factory: sessionmaker[Session], dataset_alpha: DatasetRef, context: PostingContext
) -> dict[str, int]:
    return seed_inventory_package_data(session_factory, dataset_alpha, context)


def _item(
    session_factory: sessionmaker[Session],
    dataset: DatasetRef,
    context: PostingContext,
    *,
    with_box: bool = False,
) -> tuple[int, str]:
    """Mã hàng mới (khóa tồn kho sạch); `with_box` khai thêm thùng = 12 cái."""
    code = f"HH-8C-{uuid4().hex[:6].upper()}"
    scope = posting_scope(dataset, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        item_id = create_goods_item(session, code)
        if with_box:
            session.add(ItemUnit(item_id=item_id, unit_id=UNIT_BOX_ID, factor=BOX_FACTOR))
            session.flush()
        return item_id, code


def _stock_row(
    item_code: str,
    *,
    quantity: object,
    unit_price: object,
    account: str = "156",
    warehouse: str = MAIN_CODE,
    lot_no: str | None = None,
    unit_code: str | None = None,
    amount: object = None,
    received_on: str | None = None,
    receipt_no: str | None = None,
) -> list[object]:
    return [
        account,
        warehouse,
        item_code,
        lot_no,
        unit_code,
        quantity,
        unit_price,
        amount,
        received_on,
        receipt_no,
    ]


def _opening_movements(session: Session, context: PostingContext) -> list[InventoryMovement]:
    return list(
        session.execute(
            select(InventoryMovement)
            .where(
                InventoryMovement.branch_id == context.branch_id,
                InventoryMovement.opening_layer_id.is_not(None),
            )
            .order_by(InventoryMovement.item_id, InventoryMovement.sequence_in_day)
        )
        .scalars()
        .all()
    )


def _stock_parents(session: Session, context: PostingContext) -> list[OpeningBalance]:
    return list(
        session.execute(
            select(OpeningBalance)
            .where(
                OpeningBalance.fiscal_year_id == context.fiscal_year_id,
                OpeningBalance.branch_id == context.branch_id,
                OpeningBalance.detail_kind == OpeningDetailKind.STOCK,
            )
            .order_by(OpeningBalance.item_id)
        )
        .scalars()
        .all()
    )


def _set_scope(session: Session, context: PostingContext, *, by_warehouse: bool) -> None:
    year = session.get(FiscalYear, context.fiscal_year_id)
    assert year is not None
    year.inventory_costing_by_warehouse = by_warehouse
    session.flush()


# ------------------------------------------------------------------ mẫu


def test_template_contains_stock_sheet() -> None:
    workbook = load_workbook(BytesIO(build_opening_template()))
    assert "Tồn kho" in workbook.sheetnames
    header = [cell.value for cell in next(workbook["Tồn kho"].iter_rows(max_row=1))]
    assert header[:3] == ["Số hiệu tài khoản *", "Mã kho *", "Mã hàng *"]
    assert header[5:7] == ["Số lượng *", "Đơn giá *"]


# ------------------------------------------------------------- kiểm tra cứu


def test_stock_sheet_lookup_errors(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Mỗi dòng một lỗi, cả tệp báo một lượt (trả-toàn-bộ-lỗi)."""
    _, code = _item(session_factory, dataset_alpha, context)
    report = run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of(
            {
                ACCOUNT_SHEET: [
                    ["152", None, None, None, None, 1_000, None],
                    # 158 theo dõi kho/mã hàng nhưng KHÔNG phải TK kho của gói → nhận
                    # ở sheet TK thường (review H-1), chỉ cảnh báo chiều bị mất.
                    ["158", None, None, None, None, 1_000, None],
                ],
                STOCK_SHEET: [
                    _stock_row(code, quantity=1, unit_price=100, account="632"),
                    _stock_row(code, quantity=1, unit_price=100, account="158"),
                    _stock_row(code, quantity=1, unit_price=100, warehouse="KHO-KHONG-CO"),
                    _stock_row("DV-8A-01", quantity=1, unit_price=100),
                    _stock_row("TP-8A-01", quantity=1, unit_price=100, unit_code="THUNG-8A"),
                    _stock_row(code, quantity=0, unit_price=100),
                    _stock_row(code, quantity=1, unit_price=100, received_on="2026-01-05"),
                ],
            }
        ),
        commit=False,
    )
    codes = {(error.sheet, error.row, error.code) for error in report.errors}
    assert ("Số dư tài khoản", 2, "opening.needs_stock_sheet") in codes
    assert ("Số dư tài khoản", 3, "opening.needs_stock_sheet") not in codes
    assert ("Tồn kho", 2, "opening.not_stock_account") in codes
    assert ("Tồn kho", 3, "opening.not_stock_account") in codes
    assert ("Tồn kho", 4, "opening.warehouse_unknown") in codes
    assert ("Tồn kho", 5, "opening.item_not_stocked") in codes
    assert ("Tồn kho", 6, "opening.unit_not_declared") in codes
    assert ("Tồn kho", 7, "opening.quantity_required") in codes
    assert ("Tồn kho", 8, "opening.date_not_before_year") in codes
    assert {(e.sheet, e.row) for e in report.errors if e.sheet == "Số dư tài khoản"} == {
        ("Số dư tài khoản", 2)
    }
    assert "opening.tracking_not_captured" in {w.code for w in report.warnings}
    assert not report.is_valid


def test_stock_sheet_refuses_management_ledger(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    _, code = _item(session_factory, dataset_alpha, context)
    report = run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of({STOCK_SHEET: [_stock_row(code, quantity=1, unit_price=100)]}),
        ledger=Ledger.MANAGEMENT,
        commit=False,
    )
    assert {error.code for error in report.errors} == {"opening.stock_financial_only"}


def test_fifo_year_requires_received_on(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    _, code = _item(session_factory, dataset_alpha, context)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        set_valuation_method(session, context, InventoryValuationMethod.FIFO)
    try:
        report = run_opening_import(
            session_factory,
            dataset_alpha,
            context,
            tmp_path,
            workbook_of(
                {
                    STOCK_SHEET: [
                        _stock_row(code, quantity=1, unit_price=100),
                        _stock_row(code, quantity=1, unit_price=100, received_on="2025-11-03"),
                    ]
                }
            ),
            commit=False,
        )
    finally:
        with unit_of_work(session_factory, scope) as session:
            set_valuation_method(session, context, InventoryValuationMethod.WEIGHTED_AVERAGE_MOVING)
    assert [(error.row, error.code) for error in report.errors] == [
        (2, "opening.received_on_required")
    ]


# -------------------------------------------------------------------- ghi


def test_commit_materializes_layers_as_costed_opening_movements(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """A: hai lần nhập (gõ ngược thứ tự ngày) → hai movement seq theo ngày nhập;
    B: 3 thùng × 12 = 36 cái @ 1.200/thùng → 100/cái (FR-STK-006). Cha nhóm 5 =
    Σ lớp; TK 156 lên cân đối kỳ 1 với dư Nợ đầu = 1.600 + 3.600."""
    item_a, code_a = _item(session_factory, dataset_alpha, context)
    item_b, code_b = _item(session_factory, dataset_alpha, context, with_box=True)
    report = run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of(
            {
                STOCK_SHEET: [
                    _stock_row(
                        code_a,
                        quantity=5,
                        unit_price=120,
                        received_on="01/12/2025",
                        receipt_no="NK-2",
                    ),
                    _stock_row(
                        code_a,
                        quantity=10,
                        unit_price=100,
                        received_on="2025-11-03",
                        receipt_no="NK-1",
                    ),
                    _stock_row(code_b, quantity=3, unit_price=1_200, unit_code="THUNG-8A"),
                ],
            }
        ),
    )
    assert report.is_valid and report.committed, report.errors
    assert report.rows_by_kind == {OpeningDetailKind.STOCK: 2}
    assert report.stock_layer_rows == 3
    assert report.total_debit == Decimal("5200.00")
    # 156 của gói test không theo dõi kho/mã hàng → cảnh báo, không lỗi.
    assert "opening.stock_account_untracked" in {w.code for w in report.warnings}

    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        parents = _stock_parents(session, context)
        assert [(p.item_id, p.warehouse_id, p.quantity, p.debit) for p in parents] == [
            (item_a, MAIN_WAREHOUSE_ID, Decimal("15.000000"), Decimal("1600.00")),
            (item_b, MAIN_WAREHOUSE_ID, Decimal("36.000000"), Decimal("3600.00")),
        ]
        assert [p.unit_price for p in parents] == [Decimal("106.6667"), Decimal("100.0000")]
        layers = (
            session.execute(
                select(OpeningBalanceStockLayer)
                .where(OpeningBalanceStockLayer.branch_id == context.branch_id)
                .order_by(OpeningBalanceStockLayer.sort_order)
            )
            .scalars()
            .all()
        )
        assert [(x.received_on, x.receipt_no, x.quantity, x.unit_cost) for x in layers] == [
            (DEC_1, "NK-2", Decimal(5), Decimal("120.000000")),
            (NOV_3, "NK-1", Decimal(10), Decimal("100.000000")),
            (None, None, Decimal(36), Decimal("100.000000")),
        ]
        movements = _opening_movements(session, context)
        assert [
            (m.item_id, m.sequence_in_day, m.quantity, m.unit_cost, m.amount, m.cost_state)
            for m in movements
        ] == [
            (item_a, 1, Decimal(10), Decimal("100.000000"), Decimal("1000.00"), CostState.COSTED),
            (item_a, 2, Decimal(5), Decimal("120.000000"), Decimal("600.00"), CostState.COSTED),
            (item_b, 1, Decimal(36), Decimal("100.000000"), Decimal("3600.00"), CostState.COSTED),
        ]
        assert {m.posting_date for m in movements} == {DEC_31}
        assert {m.period_id for m in movements} == {first_period_id(session, context)}
        assert all(m.voucher_id is None and m.line_id is None for m in movements)
        rows = stock_rows(session, as_of=date(2026, 1, 1), branch_id=context.branch_id)
        assert {(r.item_id, r.on_hand, r.value) for r in rows} == {
            (item_a, Decimal(15), Decimal("1600.00")),
            (item_b, Decimal(36), Decimal("3600.00")),
        }
        balance = trial_balance(
            session,
            ledger=Ledger.FINANCIAL,
            period_id=first_period_id(session, context),
            branch_id=context.branch_id,
        )
        (row_156,) = [r for r in balance.rows if r.account_code == "156"]
        assert row_156.opening_debit == Decimal("5200.00")


def test_issue_before_first_receipt_uses_opening_value_over_quantity(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """FR-STK-007 phần "chưa phát sinh nhập kho → đơn giá xuất = giá trị tồn / số
    lượng tồn": tồn đầu 15 = 1.600 → xuất 3 = round(3 × 1.600/15) = 320."""
    item, code = _item(session_factory, dataset_alpha, context)
    run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of(
            {
                STOCK_SHEET: [
                    _stock_row(code, quantity=10, unit_price=100, received_on="2025-11-03"),
                    _stock_row(code, quantity=5, unit_price=120, received_on="2025-12-01"),
                ]
            }
        ),
    )
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        issue = post_issue(
            session, context, accounts, item_id=item, posting_date=JAN_10, quantity=Decimal(3)
        )
        run_engine(session, context)
        moved = out_movement(session, issue)
        assert (moved.unit_cost, moved.amount) == (Decimal("106.666667"), Decimal("320.00"))


def test_reimport_replaces_movements_and_marks_the_year_for_recalc(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    item, code = _item(session_factory, dataset_alpha, context)
    first = workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]})
    run_opening_import(session_factory, dataset_alpha, context, tmp_path, first)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        issue = post_issue(
            session, context, accounts, item_id=item, posting_date=JAN_10, quantity=Decimal(4)
        )
        run_engine(session, context)
        assert out_movement(session, issue).amount == Decimal("400.00")
        old_ids = {m.id for m in _opening_movements(session, context)}

    second = workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=250)]})
    run_opening_import(session_factory, dataset_alpha, context, tmp_path, second)
    with unit_of_work(session_factory, scope) as session:
        movements = _opening_movements(session, context)
        assert {m.id for m in movements}.isdisjoint(old_ids)
        assert [(m.quantity, m.unit_cost) for m in movements] == [
            (Decimal(10), Decimal("250.000000"))
        ]
        marks = marks_of_branch(session, context.branch_id)
        assert [(m.item_id, m.from_date) for m in marks] == [(item, date(2026, 1, 1))]
        run_engine(session, context)
        assert out_movement(session, issue).amount == Decimal("1000.00")
        assert marks_of_branch(session, context.branch_id) == []


def test_clear_group_removes_opening_movements(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    _, code = _item(session_factory, dataset_alpha, context)
    run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]}),
    )
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        assert len(_opening_movements(session, context)) == 1
        result = run_clear_group(
            JobContext(
                job_id=uuid4(),
                session=session,
                progress=FakeProgress(reports=[]),
                attempt=1,
                dataset_schema=dataset_alpha.schema_name,
                branch_id=context.branch_id,
                requested_by=ACTOR_ID,
            ),
            ClearGroupParams(
                fiscal_year_id=context.fiscal_year_id,
                ledger=Ledger.FINANCIAL.value,
                detail_kind=OpeningDetailKind.STOCK,
            ),
        )
        assert result is not None and result["deleted_rows"] == 1
    with unit_of_work(session_factory, scope) as session:
        assert _opening_movements(session, context) == []
        assert _stock_parents(session, context) == []
        assert (
            session.scalar(
                select(OpeningBalanceStockLayer.id).where(
                    OpeningBalanceStockLayer.branch_id == context.branch_id
                )
            )
            is None
        )


def test_manual_opening_is_refused_once_the_branch_has_stock_history(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Năm 2026 nhập tay được; năm 2027 (chi nhánh đã có sổ kho 2025-12-31) thì
    không — tồn đầu 2027 là số chuyển từ 2026."""
    _, code = _item(session_factory, dataset_alpha, context)
    run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]}),
    )
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        next_year_id = ensure_fiscal_year(session, "2027", date(2027, 1, 1))
    next_context = replace(context, fiscal_year_id=next_year_id)
    with pytest.raises(OpeningStockHistoryExistsError):
        run_opening_import(
            session_factory,
            dataset_alpha,
            next_context,
            tmp_path,
            workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]}),
        )


def test_posting_before_a_manual_opening_is_refused(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Chiều ngược của bất biến (review H-3): chi nhánh mới nhập tay tồn đầu 2027
    (không có lịch sử) → ghi sổ phiếu ngày 2026 bị từ chối — engine đọc cả lịch
    sử, số ấy sẽ cộng lên tồn đầu đã khai."""
    item, code = _item(session_factory, dataset_alpha, context)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        next_year_id = ensure_fiscal_year(session, "2027", date(2027, 1, 1))
    run_opening_import(
        session_factory,
        dataset_alpha,
        replace(context, fiscal_year_id=next_year_id),
        tmp_path,
        workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]}),
    )
    with unit_of_work(session_factory, scope) as session:
        with pytest.raises(InventoryMovementBeforeOpeningStockError):
            post_receipt(
                session,
                context,
                accounts,
                item_id=item,
                posting_date=date(2026, 6, 1),
                quantity=Decimal(1),
                unit_cost=Decimal(1),
            )


def test_layer_named_by_a_specific_issue_cannot_be_replaced(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Review H-2: XK đích danh trỏ lớp đầu kỳ → nhập lại sheet Tồn kho bị guard
    chặn có tên phiếu, không phải FK RESTRICT thô."""
    item, code = _item(session_factory, dataset_alpha, context)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        set_valuation_method(session, context, InventoryValuationMethod.SPECIFIC)
    try:
        run_opening_import(
            session_factory,
            dataset_alpha,
            context,
            tmp_path,
            workbook_of(
                {
                    STOCK_SHEET: [
                        _stock_row(code, quantity=10, unit_price=100, received_on="2025-11-03")
                    ]
                }
            ),
        )
        with unit_of_work(session_factory, scope) as session:
            (layer,) = _opening_movements(session, context)
            post_issue(
                session,
                context,
                accounts,
                item_id=item,
                posting_date=JAN_10,
                quantity=Decimal(4),
                source_movement_id=layer.id,
            )
        with pytest.raises(InventoryLayerReferencedError) as refused:
            run_opening_import(
                session_factory,
                dataset_alpha,
                context,
                tmp_path,
                workbook_of(
                    {
                        STOCK_SHEET: [
                            _stock_row(code, quantity=10, unit_price=120, received_on="2025-11-03")
                        ]
                    }
                ),
            )
        assert refused.value.details["count"] == 1
    finally:
        with unit_of_work(session_factory, scope) as session:
            set_valuation_method(session, context, InventoryValuationMethod.WEIGHTED_AVERAGE_MOVING)


# ----------------------------------------------------- engine đọc lớp đầu kỳ


def test_fifo_consumes_opening_layers_in_receipt_date_order(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Sheet gõ lớp 12/2025 trước lớp 11/2025; FIFO xuất 12 = 10 × 100 + 2 × 120."""
    item, code = _item(session_factory, dataset_alpha, context)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        set_valuation_method(session, context, InventoryValuationMethod.FIFO)
    try:
        run_opening_import(
            session_factory,
            dataset_alpha,
            context,
            tmp_path,
            workbook_of(
                {
                    STOCK_SHEET: [
                        _stock_row(code, quantity=5, unit_price=120, received_on="2025-12-01"),
                        _stock_row(code, quantity=10, unit_price=100, received_on="2025-11-03"),
                    ]
                }
            ),
        )
        with unit_of_work(session_factory, scope) as session:
            issue = post_issue(
                session, context, accounts, item_id=item, posting_date=JAN_10, quantity=Decimal(12)
            )
            run_engine(session, context)
            assert out_movement(session, issue).amount == Decimal("1240.00")
    finally:
        with unit_of_work(session_factory, scope) as session:
            set_valuation_method(session, context, InventoryValuationMethod.WEIGHTED_AVERAGE_MOVING)


def test_specific_issue_can_name_an_opening_layer(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    item, code = _item(session_factory, dataset_alpha, context)
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        set_valuation_method(session, context, InventoryValuationMethod.SPECIFIC)
    try:
        run_opening_import(
            session_factory,
            dataset_alpha,
            context,
            tmp_path,
            workbook_of(
                {
                    STOCK_SHEET: [
                        _stock_row(code, quantity=10, unit_price=100, received_on="2025-11-03")
                    ]
                }
            ),
        )
        with unit_of_work(session_factory, scope) as session:
            (layer,) = _opening_movements(session, context)
            issue = post_issue(
                session,
                context,
                accounts,
                item_id=item,
                posting_date=JAN_10,
                quantity=Decimal(4),
                source_movement_id=layer.id,
            )
            run_engine(session, context)
            moved = out_movement(session, issue)
            assert (moved.unit_cost, moved.amount) == (Decimal("100.000000"), Decimal("400.00"))
    finally:
        with unit_of_work(session_factory, scope) as session:
            set_valuation_method(session, context, InventoryValuationMethod.WEIGHTED_AVERAGE_MOVING)


# ------------------------------------------------------------- chuyển năm


def test_carry_forward_writes_stock_rows_from_the_ledger_without_new_movements(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """2026: tồn đầu 10 @ 100 (156) + nhập 10 @ 200 (Nợ 156) + xuất 5 (Có 156, giá
    150) → 2027 một dòng nhóm 5 cho 156: 2.250 / 15 cái; 632 mang chiều mã hàng
    nhưng không phải TK kho → nhóm 0; không movement ngày 31/12/2026, không
    lớp; nhập tay 2027 bị từ chối (đã có sổ kho)."""
    item, code = _item(session_factory, dataset_alpha, context)
    run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of({STOCK_SHEET: [_stock_row(code, quantity=10, unit_price=100)]}),
    )
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        service = InventoryVoucherService(session)
        receipt = receipt_payload(
            context,
            accounts,
            posting_date=date(2026, 2, 5),
            quantity=Decimal(10),
            unit_cost=Decimal(200),
            item_id=item,
            operation="nhap-khac",
        )
        receipt = receipt.model_copy(
            update={
                "lines": tuple(
                    line.model_copy(
                        update={
                            "debit_account_id": accounts["156"],
                            "credit_account_id": accounts["154"],
                        }
                    )
                    for line in receipt.lines
                )
            }
        )
        voucher = service.create(receipt, user_id=ACTOR_ID)
        service.post(voucher.id, user_id=ACTOR_ID)
        issue = issue_payload(
            context, accounts, posting_date=date(2026, 2, 10), quantity=Decimal(5), item_id=item
        )
        issue = issue.model_copy(
            update={
                "operation_code": "xuat-khac",
                "lines": tuple(
                    line.model_copy(
                        update={
                            "debit_account_id": accounts["632"],
                            "credit_account_id": accounts["156"],
                        }
                    )
                    for line in issue.lines
                ),
            }
        )
        voucher = service.create(issue, user_id=ACTOR_ID)
        service.post(voucher.id, user_id=ACTOR_ID, acknowledged_warnings=True)
        run_engine(session, context)
        assert out_movement(session, voucher.id).amount == Decimal("750.00")
        next_year_id = ensure_fiscal_year(session, "2027", date(2027, 1, 1))

    with unit_of_work(session_factory, scope) as session:
        result = carry_forward(session, dataset_alpha, context)
        assert result["stock_rows_annotated"] == 2  # hai sổ, mỗi sổ một dòng 156
    with unit_of_work(session_factory, scope) as session:
        carried = list(
            session.execute(
                select(OpeningBalance)
                .where(
                    OpeningBalance.fiscal_year_id == next_year_id,
                    OpeningBalance.branch_id == context.branch_id,
                    OpeningBalance.ledger == Ledger.FINANCIAL,
                    OpeningBalance.item_id == item,
                )
                .order_by(OpeningBalance.detail_kind)
            )
            .scalars()
            .all()
        )
        by_account = {row.account_id: row for row in carried}
        stock = by_account[accounts["156"]]
        assert stock.detail_kind == OpeningDetailKind.STOCK
        assert (stock.warehouse_id, stock.debit, stock.quantity, stock.unit_price) == (
            MAIN_WAREHOUSE_ID,
            Decimal("2250.00"),
            Decimal("15.0000"),
            Decimal("150.0000"),
        )
        cogs = by_account[accounts["632"]]
        assert (cogs.detail_kind, cogs.debit, cogs.quantity) == (
            OpeningDetailKind.ACCOUNT,
            Decimal("750.00"),
            None,
        )
        assert (
            session.scalar(
                select(InventoryMovement.id).where(
                    InventoryMovement.branch_id == context.branch_id,
                    InventoryMovement.posting_date == date(2026, 12, 31),
                )
            )
            is None
        )
        assert (
            session.scalar(
                select(OpeningBalanceStockLayer.id).where(
                    OpeningBalanceStockLayer.fiscal_year_id == next_year_id,
                    OpeningBalanceStockLayer.branch_id == context.branch_id,
                )
            )
            is None
        )
    with pytest.raises(OpeningStockHistoryExistsError):
        run_opening_import(
            session_factory,
            dataset_alpha,
            replace(context, fiscal_year_id=next_year_id),
            tmp_path,
            workbook_of({STOCK_SHEET: [_stock_row(code, quantity=1, unit_price=1)]}),
        )


def test_carry_forward_leaves_quantity_empty_when_one_item_spans_two_accounts(
    session_factory: sessionmaker[Session],
    dataset_alpha: DatasetRef,
    context: PostingContext,
    accounts: dict[str, int],
    tmp_path: Path,
) -> None:
    """Review M-3: một (kho, mã hàng) trên hai TK kho (156 + 152) → 2027 hai dòng
    nhóm 5, số lượng không chia được cho từng TK → NULL cả hai, tiền vẫn đúng."""
    _, code = _item(session_factory, dataset_alpha, context)
    run_opening_import(
        session_factory,
        dataset_alpha,
        context,
        tmp_path,
        workbook_of(
            {
                STOCK_SHEET: [
                    _stock_row(code, quantity=10, unit_price=100, account="156"),
                    _stock_row(code, quantity=5, unit_price=100, account="152"),
                ]
            }
        ),
    )
    scope = posting_scope(dataset_alpha, context, user_id=ACTOR_ID)
    with unit_of_work(session_factory, scope) as session:
        next_year_id = ensure_fiscal_year(session, "2027", date(2027, 1, 1))
    with unit_of_work(session_factory, scope) as session:
        result = carry_forward(session, dataset_alpha, context)
        assert result["stock_rows_annotated"] == 0
    with unit_of_work(session_factory, scope) as session:
        rows = (
            session.execute(
                select(OpeningBalance)
                .where(
                    OpeningBalance.fiscal_year_id == next_year_id,
                    OpeningBalance.branch_id == context.branch_id,
                    OpeningBalance.ledger == Ledger.FINANCIAL,
                    OpeningBalance.detail_kind == OpeningDetailKind.STOCK,
                )
                .order_by(OpeningBalance.account_id)
            )
            .scalars()
            .all()
        )
        assert {(row.account_id, row.debit, row.quantity) for row in rows} == {
            (accounts["156"], Decimal("1000.00"), None),
            (accounts["152"], Decimal("500.00"), None),
        }
